"""Reconcile the facilities seen in ODK submissions against the national health facility registry.

Phase 0 discovery tool. Reads local files only and makes no network calls.

Writes:
  out/facility_reconciliation.xlsx  client sign-off workbook; becomes the Phase 1 seed master list
  out/facility_reconciliation.md    method, counts and the client decisions needed

Usage:
  python facility_reconcile.py --registry data/hfr_kaduna.xlsx --submissions data/submissions/
  python facility_reconcile.py --registry data/hfr.csv --map code="HF Code" --map lga="LGA Name" ...

Matching runs within one LGA only: exact code, then exact normalised name, then fuzzy
(rapidfuzz token_set_ratio). See discovery/README.md for the thresholds and guards.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from rapidfuzz import fuzz, process
from rich.console import Console
from rich.table import Table

HERE = Path(__file__).resolve().parent
OUT_DIR = HERE / "out"
XLSX_OUT = OUT_DIR / "facility_reconciliation.xlsx"
MD_OUT = OUT_DIR / "facility_reconciliation.md"

AUTO_MATCH = 95          # token_set_ratio at or above this may auto-match
REVIEW_MIN = 60          # below this is Unmatched
AMBIGUITY_MARGIN = 3     # runner-up this close to the top score blocks an auto-match
CORE_MIN = 90            # facility-type-stripped names must agree this well to auto-match
EXPECTED_TOTAL = 2200
DECISION_DAYS = 10       # business days from the run date to the suggested deadline
DECISIONS = ("MAP", "ADD", "REJECT")

console = Console(stderr=True)

# --------------------------------------------------------------------------- LGAs

LGAS = [
    "Birnin Gwari", "Chikun", "Giwa", "Igabi", "Ikara", "Jaba", "Jema'a", "Kachia",
    "Kaduna North", "Kaduna South", "Kagarko", "Kajuru", "Kaura", "Kauru", "Kubau", "Kudan",
    "Lere", "Makarfi", "Sabon Gari", "Sanga", "Soba", "Zangon Kataf", "Zaria",
]
assert len(LGAS) == 23


def compact(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", value.lower())


# Spellings seen in Nigerian registries, keyed by compact form. Exact lookups only: fuzzy LGA
# matching could pull in look-alike LGAs from other states (e.g. Kaura Namoda, Zamfara).
_LGA_INDEX: dict[str, str] = {compact(n): n for n in LGAS} | {
    "zangokataf": "Zangon Kataf",
    "zangonkatab": "Zangon Kataf",
    "jemaa": "Jema'a",
    "jemaah": "Jema'a",
    "sabongarizaria": "Sabon Gari",
    "birnigwari": "Birnin Gwari",
    "kadunan": "Kaduna North",
    "kadunas": "Kaduna South",
    "kadunanth": "Kaduna North",
    "kadunasth": "Kaduna South",
}


def resolve_lga(raw: str, aliases: dict[str, str]) -> str | None:
    key = compact(raw)
    if not key:
        return None
    return aliases.get(key) or _LGA_INDEX.get(key)


# --------------------------------------------------------------------------- normalisation

# Token-level abbreviations, applied after punctuation is stripped.
_TOKEN_ABBREV = {
    "phcc": "phc", "comp": "comprehensive", "compr": "comprehensive",
    "ctr": "centre", "cntr": "centre", "cent": "centre", "center": "centre", "centres": "centre",
    "centers": "centre", "hosp": "hospital", "gen": "general", "clin": "clinic", "cln": "clinic",
    "clinics": "clinic", "hlth": "health", "pry": "primary", "govt": "government",
    "gov": "government", "matern": "maternity", "disp": "dispensary",
}

# Phrases collapsed to one canonical token, longest first.
_PHRASES = [
    (r"\bprimary health care centre\b", "phc"),
    (r"\bprimary health centre\b", "phc"),
    (r"\bprimary health care\b", "phc"),
    (r"\bcomprehensive health centre\b", "chc"),
    (r"\bmaternal (and )?child health\b", "mch"),
    (r"\bhealth centre\b", "hc"),
    (r"\bhealth post\b", "hp"),
]

# Words that describe the type of facility rather than which one it is.
_TYPE_TOKENS = {
    "phc", "chc", "hc", "hp", "mch", "clinic", "centre", "comprehensive", "primary", "health",
    "care", "post", "hospital", "general", "maternity", "dispensary", "cottage", "and", "of", "the",
}


def normalise_name(raw: str) -> str:
    """Canonical form of a facility name. Empty when the value is not a name (e.g. "2019.00")."""
    s = raw.lower().replace("&", " and ")
    s = re.sub(r"['`’]", "", s)
    s = s.replace("healthcare", "health care")
    s = re.sub(r"[^a-z0-9]+", " ", s)
    # "p h c" (from "P.H.C.") -> "phc"
    s = re.sub(r"\b(?:[a-z] ){1,}[a-z]\b", lambda m: m.group(0).replace(" ", ""), s)
    tokens = [_TOKEN_ABBREV.get(t, t) for t in s.split()]
    s = " ".join(tokens)
    for pattern, repl in _PHRASES:
        s = re.sub(pattern, repl, s)
    s = re.sub(r"\s+", " ", s).strip()
    return s if re.search(r"[a-z]", s) else ""


def core_name(norm: str) -> str:
    return " ".join(t for t in norm.split() if t not in _TYPE_TOKENS)


def normalise_code(raw: str) -> str:
    return re.sub(r"\s+", "", raw).upper()


_OWNERSHIP = {
    "public": "Public", "government": "Public", "govt": "Public", "public government": "Public",
    "private": "Private", "private for profit": "Private",
}
_LEVEL = {
    "primary": "Primary", "primary health care": "Primary", "phc": "Primary",
    "secondary": "Secondary", "tertiary": "Tertiary",
}


def classify(raw: str, table: dict[str, str]) -> str:
    """Public/Private or Primary/Secondary/Tertiary; "" when the value needs a client decision."""
    return table.get(re.sub(r"[^a-z]+", " ", raw.lower()).strip(), "")


# --------------------------------------------------------------------------- table input


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def read_table(path: Path, sheet: str | None = None) -> tuple[list[str], list[list[str]]]:
    """Header row and data rows of a CSV or XLSX file, all values as stripped strings."""
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        wb = load_workbook(path, read_only=True, data_only=True)
        if sheet and sheet not in wb.sheetnames:
            sys.exit(f"{path.name} has no sheet {sheet!r}. Sheets: {', '.join(wb.sheetnames)}")
        ws = wb[sheet] if sheet else wb.worksheets[0]
        if not sheet and len(wb.sheetnames) > 1:
            console.print(f"[yellow]{path.name}: reading first sheet {ws.title!r} "
                          f"(use --registry-sheet to choose)[/]")
        raw_rows = [[cell_text(v) for v in row] for row in ws.iter_rows(values_only=True)]
        wb.close()
    elif path.suffix.lower() in (".csv", ".txt"):
        text = _read_text(path)
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        raw_rows = [[c.strip() for c in row] for row in csv.reader(text.splitlines(), dialect)]
    else:
        sys.exit(f"Unsupported file type: {path} (expected .csv or .xlsx)")

    # Header = first row with at least three non-empty cells (skips export title rows).
    for i, row in enumerate(raw_rows[:20]):
        if sum(1 for c in row if c) >= 3:
            headers = row
            body = [r for r in raw_rows[i + 1:] if any(r)]
            return headers, [r + [""] * (len(headers) - len(r)) for r in body]
    sys.exit(f"Could not find a header row in the first 20 rows of {path}")


def _read_text(path: Path) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    sys.exit(f"Could not decode {path} as UTF-8 or Windows-1252")


def _norm_header(h: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", h.lower()).strip()


def parse_map(pairs: list[str], allowed: set[str], option: str) -> dict[str, str]:
    mapping: dict[str, str] = {}
    for pair in pairs:
        key, sep, value = pair.partition("=")
        key = key.strip().lower()
        if not sep or key not in allowed or not value.strip():
            sys.exit(f"{option} {pair!r}: expected FIELD=COLUMN with FIELD one of "
                     f"{', '.join(sorted(allowed))}")
        mapping[key] = value.strip()
    return mapping


# --------------------------------------------------------------------------- registry

REGISTRY_FIELDS: dict[str, list[str]] = {
    "code": ["facility code", "hf code", "hfr code", "facility uid", "uid", "facility id",
             "code", "unique id", "facility unique id"],
    "name": ["facility name", "name", "health facility name", "hf name", "facility",
             "name of facility"],
    "lga": ["lga", "lga name", "local government area", "local government", "lga lga"],
    "ward": ["ward", "ward name"],
    "ownership": ["ownership", "owner", "facility ownership", "ownership type"],
    "level": ["facility level", "level", "level of care", "facility level of care"],
    "state": ["state", "state name"],
}
REGISTRY_REQUIRED = ("name", "lga")


@dataclass
class RegistryFacility:
    row: int
    code: str
    name: str
    lga: str
    ward: str
    ownership_raw: str
    level_raw: str
    norm: str = ""
    core: str = ""
    ownership: str = ""
    level: str = ""

    def __post_init__(self) -> None:
        self.norm = normalise_name(self.name)
        self.core = core_name(self.norm)
        self.ownership = classify(self.ownership_raw, _OWNERSHIP)
        self.level = classify(self.level_raw, _LEVEL)


def resolve_registry_columns(headers: list[str], overrides: dict[str, str],
                             path: Path) -> dict[str, int]:
    index: dict[str, int] = {}
    ambiguous: dict[str, list[str]] = {}
    lower = {h.lower(): i for i, h in enumerate(headers)}
    for fld, synonyms in REGISTRY_FIELDS.items():
        if fld in overrides:
            i = lower.get(overrides[fld].lower())
            if i is None:
                sys.exit(f"--map {fld}={overrides[fld]!r}: no such column in {path.name}. "
                         f"Columns: {', '.join(repr(h) for h in headers if h)}")
            index[fld] = i
            continue
        hits = [i for i, h in enumerate(headers) if _norm_header(h) in synonyms]
        if len(hits) == 1:
            index[fld] = hits[0]
        elif len(hits) > 1:
            ambiguous[fld] = [headers[i] for i in hits]

    missing = [f for f in REGISTRY_REQUIRED if f not in index]
    if missing or ambiguous:
        lines = [f"Cannot map the registry columns in {path.name}; please give a column mapping."]
        if missing:
            lines.append(f"  Not found: {', '.join(missing)}")
        for fld, cols in ambiguous.items():
            lines.append(f"  Ambiguous {fld}: {', '.join(repr(c) for c in cols)}")
        lines.append(f"  Columns: {', '.join(repr(h) for h in headers if h)}")
        lines.append('  Example: --map code="Facility Code" --map name="Facility Name" '
                     '--map lga="LGA" --map ward="Ward" --map ownership="Ownership" '
                     '--map level="Facility Level"')
        sys.exit("\n".join(lines))
    return index


@dataclass
class RegistryLoad:
    facilities: list[RegistryFacility]
    columns: dict[str, str]
    rows_read: int
    other_state: int
    unknown_lga: Counter
    duplicate_codes: int


def load_registry(path: Path, sheet: str | None, overrides: dict[str, str],
                  aliases: dict[str, str]) -> RegistryLoad:
    headers, rows = read_table(path, sheet)
    idx = resolve_registry_columns(headers, overrides, path)
    get = lambda r, f: r[idx[f]] if f in idx else ""  # noqa: E731

    facilities: list[RegistryFacility] = []
    other_state = 0
    unknown_lga: Counter = Counter()
    seen_codes: set[str] = set()
    duplicate_codes = 0
    for n, r in enumerate(rows, start=2):
        if "state" in idx and "kaduna" not in get(r, "state").lower():
            other_state += 1
            continue
        lga = resolve_lga(get(r, "lga"), aliases)
        if lga is None:
            unknown_lga[get(r, "lga") or "(blank)"] += 1
            continue
        code = normalise_code(get(r, "code"))
        if code and code in seen_codes:
            duplicate_codes += 1
        seen_codes.add(code)
        facilities.append(RegistryFacility(
            row=n, code=code, name=get(r, "name"), lga=lga, ward=get(r, "ward"),
            ownership_raw=get(r, "ownership"), level_raw=get(r, "level"),
        ))
    return RegistryLoad(
        facilities=facilities,
        columns={f: headers[i] for f, i in idx.items()},
        rows_read=len(rows),
        other_state=other_state,
        unknown_lga=unknown_lga,
        duplicate_codes=duplicate_codes,
    )


# --------------------------------------------------------------------------- submissions

# Leaf question names tried in order when --sub-map does not give a path. Form versions drift,
# so resolution is per record.
SUBMISSION_FIELDS: dict[str, list[str]] = {
    "lga": ["lga", "lga_name", "lga_code", "local_government", "local_govt"],
    "ward": ["ward", "ward_name", "ward_code"],
    "code": ["facility_code", "fac_code", "fac_id", "hf_code", "facility_id", "facility_uid"],
    "name": ["facility_name", "fac_name", "hf_name", "health_facility", "health_facility_name",
             "name_of_facility", "facility"],
    "ownership": ["ownership", "facility_ownership", "owner", "owner_type"],
    "level": ["facility_level", "level", "level_of_care"],
}
_SUBMITTED_AT = ["__system/submissionDate", "SubmissionDate", "submissionDate"]


def flatten(obj: Any, prefix: str = "") -> dict[str, str]:
    """Nested OData record -> {"group/question": value}. Repeat groups (lists) are skipped."""
    out: dict[str, str] = {}
    if isinstance(obj, dict):
        for k, v in obj.items():
            path = f"{prefix}/{k}" if prefix else str(k)
            if isinstance(v, dict):
                out.update(flatten(v, path))
            elif not isinstance(v, list):
                out[path] = cell_text(v)
    return out


def iter_submission_records(source: Path) -> list[dict[str, str]]:
    files = sorted(p for p in source.rglob("*") if p.suffix.lower() in (".json", ".csv", ".xlsx")
                   ) if source.is_dir() else [source]
    if not files:
        sys.exit(f"No .json, .csv or .xlsx files under {source}")
    records: list[dict[str, str]] = []
    for f in files:
        if f.suffix.lower() == ".json":
            data = json.loads(_read_text(f))
            items = data.get("value", []) if isinstance(data, dict) else data
            records.extend(flatten(item) for item in items if isinstance(item, dict))
        else:
            # ODK Central CSV exports join group names with "-".
            headers, rows = read_table(f)
            keys = [h.replace("-", "/") for h in headers]
            records.extend(dict(zip(keys, row)) for row in rows)
    return records


@dataclass
class SubmittedFacility:
    ref: str
    lga: str | None
    lga_raw: str
    code: str
    code_in_registry: bool
    names: Counter = field(default_factory=Counter)
    wards: Counter = field(default_factory=Counter)
    ownership: Counter = field(default_factory=Counter)
    level: Counter = field(default_factory=Counter)
    count: int = 0
    first_seen: str = ""
    last_seen: str = ""
    status: str = ""                 # matched | review | unmatched
    method: str = ""
    score: float | None = None
    reason: str = ""
    match: RegistryFacility | None = None
    candidates: list[tuple[RegistryFacility, float]] = field(default_factory=list)

    @property
    def name(self) -> str:
        return self.names.most_common(1)[0][0] if self.names else ""

    @property
    def norms(self) -> list[str]:
        return list(dict.fromkeys(n for n in (normalise_name(x) for x in self.names) if n))

    def top(self, c: Counter) -> str:
        return c.most_common(1)[0][0] if c else ""


@dataclass
class SubmissionLoad:
    facilities: list[SubmittedFacility]
    records: int
    skipped: int
    paths_used: dict[str, Counter]


def resolve_field(flat: dict[str, str], leaves: dict[str, str], fld: str,
                  sub_map: dict[str, str]) -> tuple[str, str]:
    if fld in sub_map:
        path = sub_map[fld]
        return path, flat.get(path, flat.get(path.replace("-", "/"), ""))
    for synonym in SUBMISSION_FIELDS[fld]:
        if synonym in leaves:
            path = leaves[synonym]
            return path, flat[path]
    return "", ""


def load_submissions(source: Path, sub_map: dict[str, str], aliases: dict[str, str],
                     registry_codes: set[str]) -> SubmissionLoad:
    records = iter_submission_records(source)
    groups: dict[tuple[str, str], SubmittedFacility] = {}
    paths_used: dict[str, Counter] = defaultdict(Counter)
    skipped = 0
    all_leaves: Counter = Counter()

    for flat in records:
        leaves: dict[str, str] = {}
        for path in flat:
            leaves.setdefault(path.rsplit("/", 1)[-1].lower(), path)
        all_leaves.update(flat.keys())
        values: dict[str, str] = {}
        for fld in SUBMISSION_FIELDS:
            path, value = resolve_field(flat, leaves, fld, sub_map)
            values[fld] = value
            if path:
                paths_used[fld][path] += 1
        if not values["lga"] or not (values["name"] or values["code"]):
            skipped += 1
            continue

        lga = resolve_lga(values["lga"], aliases)
        code = normalise_code(values["code"])
        # Group by code only when the registry knows it; a junk code shared by many
        # facilities would otherwise merge them.
        in_registry = bool(code) and code in registry_codes
        lga_key = lga or f"?{compact(values['lga'])}"
        key = (lga_key, f"code:{code}" if in_registry
               else f"name:{normalise_name(values['name']) or values['name'].lower()}")
        sub = groups.get(key)
        if sub is None:
            sub = groups[key] = SubmittedFacility(
                ref="", lga=lga, lga_raw=values["lga"], code=code, code_in_registry=in_registry)
        if values["name"]:
            sub.names[values["name"]] += 1
        for fld in ("ward", "ownership", "level"):
            if values[fld]:
                getattr(sub, fld if fld != "ward" else "wards")[values[fld]] += 1
        if code and not sub.code:
            sub.code = code
        sub.count += 1
        submitted = next((flat[k] for k in _SUBMITTED_AT if flat.get(k)), "")[:10]
        if submitted:
            sub.first_seen = min(filter(None, [sub.first_seen, submitted]))
            sub.last_seen = max(sub.last_seen, submitted)

    if records and skipped == len(records):
        sample = ", ".join(p for p, _ in all_leaves.most_common(60))
        sys.exit("No submission had a recognisable LGA plus facility name or code; please give "
                 "the question paths with --sub-map, e.g. --sub-map lga=grp_fac/lga "
                 f"--sub-map name=grp_fac/fac_name --sub-map code=grp_fac/fac_id\n"
                 f"  Paths seen: {sample}")

    facilities = sorted(groups.values(), key=lambda s: (s.lga or "~" + s.lga_raw, s.name, s.code))
    for i, sub in enumerate(facilities, start=1):
        sub.ref = f"S-{i:04d}"
    return SubmissionLoad(facilities, len(records), skipped, dict(paths_used))


# --------------------------------------------------------------------------- matching


def match_all(subs: list[SubmittedFacility], registry: list[RegistryFacility]) -> None:
    by_lga: dict[str, list[RegistryFacility]] = defaultdict(list)
    code_home: dict[str, str] = {}
    for r in registry:
        by_lga[r.lga].append(r)
        if r.code:
            code_home.setdefault(r.code, r.lga)
    for sub in subs:
        match_one(sub, by_lga.get(sub.lga or "", []), code_home)


def match_one(sub: SubmittedFacility, pool: list[RegistryFacility],
              code_home: dict[str, str]) -> None:
    notes: list[str] = []

    def done(status: str, reason: str = "", method: str = "",
             match: RegistryFacility | None = None, score: float | None = None) -> None:
        sub.status, sub.method, sub.match, sub.score = status, method, match, score
        sub.reason = "; ".join(filter(None, [reason, *notes]))

    if sub.lga is None:
        return done("unmatched", f"LGA {sub.lga_raw!r} is not one of the 23")
    if not pool:
        return done("unmatched", "the registry has no facilities in this LGA")

    # 1. Exact code within the LGA.
    if sub.code:
        hit = next((r for r in pool if r.code == sub.code), None)
        if hit:
            sub.candidates = [(hit, 100.0)]
            return done("matched", method="code", match=hit, score=100.0)
        if sub.code in code_home:
            notes.append(f"code {sub.code} is registered under {code_home[sub.code]}")

    norms = sub.norms
    if not norms:
        return done("unmatched", "no usable facility name, and no registry code in this LGA")

    # 2. Exact normalised name.
    exact = [r for r in pool if r.norm and r.norm in norms]
    if len(exact) == 1:
        sub.candidates = [(exact[0], 100.0)]
        return done("matched", method="exact name", match=exact[0], score=100.0)
    if len(exact) > 1:
        sub.candidates = [(r, 100.0) for r in exact[:3]]
        return done("review", f"{len(exact)} registry facilities in this LGA share this name",
                    score=100.0)

    # 3. Fuzzy, best score per registry facility across the submitted name variants.
    choices = {i: r.norm for i, r in enumerate(pool) if r.norm}
    best: dict[int, float] = {}
    for norm in norms:
        for _, score, i in process.extract(norm, choices, scorer=fuzz.token_set_ratio, limit=5):
            best[i] = max(best.get(i, 0.0), score)
    ranked = sorted(best.items(), key=lambda kv: -kv[1])
    sub.candidates = [(pool[i], s) for i, s in ranked[:3]]
    if not sub.candidates:
        return done("unmatched", "no registry facility in this LGA has a usable name")

    top, score = sub.candidates[0]
    if score < REVIEW_MIN:
        return done("unmatched", f"best score {score:.0f} is below {REVIEW_MIN}", score=score)
    if score < AUTO_MATCH:
        return done("review", f"score {score:.0f} is between {REVIEW_MIN} and {AUTO_MATCH - 1}",
                    score=score)

    # token_set_ratio scores 100 whenever one name's words are a subset of the other's
    # ("PHC" vs "PHC Kawo"), so an auto-match also needs a clear winner and agreeing core names.
    if len(sub.candidates) > 1 and score - sub.candidates[1][1] < AMBIGUITY_MARGIN:
        return done("review", "two registry facilities score almost the same", score=score)
    cores = [core_name(n) for n in norms]
    core_score = max((fuzz.token_sort_ratio(c, top.core) for c in cores if c and top.core),
                     default=0.0)
    if core_score < CORE_MIN:
        return done("review", "names agree only on facility-type words, or one has extra words",
                    score=score)
    return done("matched", method="fuzzy", match=top, score=score)


# --------------------------------------------------------------------------- workbook

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
BOLD = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")


def write_rows(ws, headers: list[str], rows: list[list[Any]], widths: list[int],
               freeze: str = "A2", input_cols: tuple[str, ...] = ()) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 32
    for row in rows:
        ws.append(row)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = freeze
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    for name in input_cols:
        col = get_column_letter(headers.index(name) + 1)
        for r in range(2, len(rows) + 2):
            ws[f"{col}{r}"].fill = INPUT_FILL


def add_decision_validation(ws, headers: list[str], n_rows: int) -> None:
    if not n_rows:
        return
    col = get_column_letter(headers.index("Decision") + 1)
    dv = DataValidation(type="list", formula1=f'"{",".join(DECISIONS)}"', allow_blank=True)
    dv.error = "Choose MAP, ADD or REJECT."
    dv.errorTitle = "Decision"
    dv.prompt = ("MAP = same facility as a registry entry (enter its code). "
                 "ADD = real facility missing from the registry. REJECT = not a real facility.")
    dv.promptTitle = "Decision"
    dv.showErrorMessage = dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}{n_rows + 1}")


def join_counter(c: Counter) -> str:
    return " | ".join(v for v, _ in c.most_common())


def ward_agrees(sub: SubmittedFacility, reg: RegistryFacility) -> str:
    if not sub.wards or not reg.ward:
        return ""
    return "Yes" if any(fuzz.ratio(compact(w), compact(reg.ward)) >= 85 for w in sub.wards) else "No"


def build_workbook(ctx: "Run") -> Workbook:
    wb = Workbook()
    readme = wb.active
    readme.title = "Read me"

    matched = [s for s in ctx.subs if s.status == "matched"]
    review = [s for s in ctx.subs if s.status == "review"]
    unmatched = [s for s in ctx.subs if s.status == "unmatched"]

    mapped_by: dict[int, list[str]] = defaultdict(list)
    for s in matched:
        mapped_by[id(s.match)].append(s.ref)
    candidate_for: dict[int, list[str]] = defaultdict(list)
    for s in review:
        for r, _ in s.candidates:
            candidate_for[id(r)].append(s.ref)

    # Matched
    ws = wb.create_sheet("Matched")
    headers = ["Ref", "LGA", "Submitted code", "Submitted name(s)", "Submitted ward(s)",
               "Submissions", "Registry code", "Registry name", "Registry ward", "Level",
               "Ownership", "Method", "Score", "Ward agrees", "Other refs mapped here", "Note"]
    rows = []
    for s in matched:
        r = s.match
        assert r is not None
        others = [x for x in mapped_by[id(r)] if x != s.ref]
        rows.append([s.ref, s.lga, s.code, join_counter(s.names), join_counter(s.wards), s.count,
                     r.code, r.name, r.ward, r.level or r.level_raw,
                     r.ownership or r.ownership_raw, s.method, round(s.score or 0, 1),
                     ward_agrees(s, r), ", ".join(others), s.reason])
    write_rows(ws, headers, rows, [8, 14, 14, 36, 18, 11, 14, 36, 18, 11, 11, 11, 7, 8, 14, 30],
               freeze="C2")

    # Review
    ws = wb.create_sheet("Review")
    headers = ["Ref", "LGA", "Submitted code", "Submitted name(s)", "Submitted ward(s)",
               "Ownership (submitted)", "Level (submitted)", "Submissions", "Why review"]
    for n in (1, 2, 3):
        headers += [f"Candidate {n} code", f"Candidate {n} name", f"Candidate {n} ward",
                    f"Candidate {n} score"]
    headers += ["Decision", "Map to registry code", "Notes"]
    rows = []
    for s in review:
        row = [s.ref, s.lga, s.code, join_counter(s.names), join_counter(s.wards),
               s.top(s.ownership), s.top(s.level), s.count, s.reason]
        for n in range(3):
            if n < len(s.candidates):
                r, score = s.candidates[n]
                row += [r.code, r.name, r.ward, round(score, 1)]
            else:
                row += ["", "", "", ""]
        rows.append(row + ["", "", ""])
    write_rows(ws, headers, rows,
               [8, 14, 14, 34, 16, 12, 11, 11, 30] + [13, 30, 14, 9] * 3 + [11, 16, 30],
               freeze="E2", input_cols=("Decision", "Map to registry code", "Notes"))
    add_decision_validation(ws, headers, len(rows))

    # Unmatched
    ws = wb.create_sheet("Unmatched")
    headers = ["Ref", "LGA", "LGA as submitted", "Submitted code", "Submitted name(s)",
               "Submitted ward(s)", "Ownership (submitted)", "Level (submitted)", "Submissions",
               "First submitted", "Last submitted", "Why unmatched", "Best candidate code",
               "Best candidate name", "Best score", "Decision", "Map to registry code", "Notes"]
    rows = []
    for s in unmatched:
        best = s.candidates[0] if s.candidates else None
        rows.append([s.ref, s.lga or "", s.lga_raw, s.code, join_counter(s.names),
                     join_counter(s.wards), s.top(s.ownership), s.top(s.level), s.count,
                     s.first_seen, s.last_seen, s.reason,
                     best[0].code if best else "", best[0].name if best else "",
                     round(best[1], 1) if best else "", "", "", ""])
    write_rows(ws, headers, rows,
               [8, 14, 14, 14, 34, 16, 12, 11, 11, 12, 12, 34, 13, 30, 8, 11, 16, 30],
               freeze="E2", input_cols=("Decision", "Map to registry code", "Notes"))
    add_decision_validation(ws, headers, len(rows))

    # Registry never assessed
    ws = wb.create_sheet("Registry never assessed")
    headers = ["LGA", "Ward", "Registry code", "Registry name", "Level", "Ownership",
               "Review candidate for", "Registry row"]
    rows = [[r.lga, r.ward, r.code, r.name, r.level or r.level_raw,
             r.ownership or r.ownership_raw, ", ".join(candidate_for[id(r)]), r.row]
            for r in sorted(ctx.never_assessed, key=lambda r: (r.lga, r.ward, r.name))]
    write_rows(ws, headers, rows, [14, 18, 14, 36, 11, 11, 20, 10])

    build_summary(wb.create_sheet("Summary"), ctx)
    build_readme(readme, ctx)
    return wb


def per_lga_counts(ctx: "Run") -> list[list[Any]]:
    rows = []
    for lga in LGAS:
        reg = [r for r in ctx.registry.facilities if r.lga == lga]
        subs = [s for s in ctx.subs if s.lga == lga]
        rows.append([
            lga, len(reg),
            sum(r.ownership == "Public" for r in reg), sum(r.ownership == "Private" for r in reg),
            sum(not r.ownership for r in reg), len(subs),
            sum(s.status == "matched" for s in subs), sum(s.status == "review" for s in subs),
            sum(s.status == "unmatched" for s in subs),
            sum(r.lga == lga for r in ctx.never_assessed),
        ])
    unknown = [s for s in ctx.subs if s.lga is None]
    if unknown:
        rows.append(["(LGA not recognised)", 0, 0, 0, 0, len(unknown), 0, 0, len(unknown), 0])
    return rows


LGA_HEADERS = ["LGA", "Registry facilities", "Public", "Private", "Ownership unclassified",
               "Submitted facilities", "Matched", "Review", "Unmatched",
               "Registry never assessed"]


def ownership_counts(ctx: "Run") -> list[list[Any]]:
    assessed = ctx.assessed_registry
    rows = []
    for label, key in (("Public", "Public"), ("Private", "Private"), ("Unclassified", "")):
        rows.append([label,
                     sum(r.ownership == key for r in ctx.registry.facilities),
                     sum(r.ownership == key for r in assessed),
                     sum(r.ownership == key for r in ctx.never_assessed)])
    return rows


def build_summary(ws, ctx: "Run") -> None:
    lga_rows = per_lga_counts(ctx)
    totals = ["Total"] + [sum(r[i] for r in lga_rows) for i in range(1, len(LGA_HEADERS))]
    write_rows(ws, LGA_HEADERS, lga_rows + [totals], [22, 12, 9, 9, 13, 12, 10, 10, 11, 13])
    ws.auto_filter.ref = None
    for cell in ws[ws.max_row]:
        cell.font = BOLD

    ws.append([])
    ws.append(["Ownership (registry classification)", "Registry", "Assessed (matched)",
               "Never assessed"])
    for cell in ws[ws.max_row]:
        cell.font = BOLD
    for row in ownership_counts(ctx):
        ws.append(row)

    ws.append([])
    ws.append(["Totals", "Count"])
    for cell in ws[ws.max_row]:
        cell.font = BOLD
    for row in total_rows(ctx):
        ws.append(row)


def total_rows(ctx: "Run") -> list[list[Any]]:
    n_reg = len(ctx.registry.facilities)
    return [
        ["Expected facilities (client estimate)", f"≈ {EXPECTED_TOTAL:,}"],
        ["Registry facilities in the 23 LGAs", n_reg],
        ["Difference from expected", n_reg - EXPECTED_TOTAL],
        ["Distinct submitted facilities (before review)", len(ctx.subs)],
        ["Registry facilities matched at least once", len(ctx.assessed_registry)],
        ["Registry facilities never assessed", len(ctx.never_assessed)],
    ]


def build_readme(ws, ctx: "Run") -> None:
    lines: list[tuple[str, str]] = [
        ("Kaduna eDQA: facility reconciliation", ""),
        ("", ""),
        ("Purpose", "Every facility that appears in the ODK submissions, matched to the national "
                    "health facility registry. Once signed off, this becomes the facility master "
                    "list for the new portal."),
        ("Please decide by", ctx.deadline.strftime("%A %d %B %Y")),
        ("", ""),
        ("Sheet", "What to do"),
        ("Matched", "Matched automatically (same code, same name, or name score 95 or more). "
                    "Check a sample; tell us about any row that is wrong."),
        ("Review", "Possible matches (score 60 to 94, or a high score we could not confirm). "
                   "Choose a Decision for every row."),
        ("Unmatched", "No good match in the registry. Choose a Decision for every row."),
        ("Registry never assessed", "Registry facilities in Kaduna that no submission matched. "
                                    "Tell us which are closed, and which should be assessed."),
        ("Summary", "Counts per LGA, public vs private, and totals."),
        ("", ""),
        ("Decision", "Meaning"),
        ("MAP", "Same facility as a registry entry. Put its code in 'Map to registry code'."),
        ("ADD", "A real facility missing from the registry. It will be added to the master list."),
        ("REJECT", "Not a real facility (test entry, typo, duplicate). Its submissions will not "
                   "be scored."),
        ("", ""),
        ("Only edit the yellow columns. Do not sort away or delete rows; the Ref column links "
         "each row back to the submissions.", ""),
        ("", ""),
        ("Run details", ""),
        ("Generated", ctx.generated.strftime("%Y-%m-%d %H:%M")),
        ("Registry file", ctx.registry_path.name),
        ("Registry columns used", ", ".join(f"{k}={v}" for k, v in ctx.registry.columns.items())),
        ("Submissions source", ctx.submissions_path.name),
        ("Thresholds", f"auto-match ≥ {AUTO_MATCH}, review {REVIEW_MIN}–{AUTO_MATCH - 1}, "
                       f"unmatched < {REVIEW_MIN} (rapidfuzz token_set_ratio)"),
    ]
    for a, b in lines:
        ws.append([a, b])
    ws["A1"].font = Font(bold=True, size=14)
    for row in ws.iter_rows(min_row=2):
        if row[0].value in ("Purpose", "Please decide by", "Sheet", "Decision", "Run details"):
            row[0].font = row[1].font = BOLD
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 100
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = WRAP


# --------------------------------------------------------------------------- report


def add_business_days(start: date, days: int) -> date:
    d = start
    while days:
        d += timedelta(days=1)
        if d.weekday() < 5:
            days -= 1
    return d


def md_table(headers: list[str], rows: list[list[Any]]) -> str:
    def fmt(v: Any) -> str:
        return f"{v:,}" if isinstance(v, int) else str(v).replace("|", "\\|")
    out = ["| " + " | ".join(headers) + " |", "|" + "---|" * len(headers)]
    out += ["| " + " | ".join(fmt(v) for v in row) + " |" for row in rows]
    return "\n".join(out)


def build_markdown(ctx: "Run") -> str:
    reg = ctx.registry
    status = Counter(s.status for s in ctx.subs)
    methods = Counter(s.method for s in ctx.subs if s.status == "matched")
    review_reasons = Counter(s.reason.split(";")[0] for s in ctx.subs if s.status == "review")
    unmatched_reasons = Counter(
        re.sub(r"LGA '.*' is", "LGA … is", s.reason.split(";")[0])
        for s in ctx.subs if s.status == "unmatched")
    unknown_sub_lgas = Counter(s.lga_raw for s in ctx.subs if s.lga is None)
    unclassified_owner = Counter(r.ownership_raw or "(blank)" for r in reg.facilities
                                 if not r.ownership)
    unclassified_level = Counter(r.level_raw or "(blank)" for r in reg.facilities if not r.level)
    without_code = sum(1 for r in reg.facilities if not r.code)
    lga_rows = per_lga_counts(ctx)
    totals = ["**Total**"] + [sum(r[i] for r in lga_rows) for i in range(1, len(LGA_HEADERS))]

    def counts(c: Counter) -> str:
        return ", ".join(f"`{k}` ({v})" for k, v in c.most_common(15)) or "none"

    decisions: list[str] = []
    n = 0

    def decide(text: str) -> None:
        nonlocal n
        n += 1
        decisions.append(f"{n}. {text}")

    decide(f"**Sign-off owner.** Who in the client team signs this list off (open question "
           f"Q-06)? The same person should own changes to the facility list after launch.")
    if status["review"]:
        decide(f"**Review sheet:** a MAP / ADD / REJECT decision for each of "
               f"{status['review']:,} possible matches.")
    if status["unmatched"]:
        decide(f"**Unmatched sheet:** a decision for each of {status['unmatched']:,} submitted "
               f"facilities with no good registry match.")
    if ctx.never_assessed:
        decide(f"**Registry never assessed:** {len(ctx.never_assessed):,} Kaduna registry "
               f"facilities have no matching submission. Which are closed or out of scope, and "
               f"which should be assessed from now on?")
    if unknown_sub_lgas:
        decide(f"**Unrecognised LGA values in submissions:** {counts(unknown_sub_lgas)}. "
               f"Which of the 23 LGAs does each mean?")
    if reg.unknown_lga and "state" in reg.columns:
        decide(f"**Unrecognised LGA values in the registry:** {counts(reg.unknown_lga)} "
               f"({sum(reg.unknown_lga.values()):,} rows excluded).")
    if unclassified_owner:
        decide(f"**Ownership values that are not clearly Public or Private:** "
               f"{counts(unclassified_owner)}. The portal reports Public vs Private only; "
               f"where does each belong (e.g. faith-based, NGO)?")
    if unclassified_level:
        decide(f"**Level values that are not Primary / Secondary / Tertiary:** "
               f"{counts(unclassified_level)}.")
    if without_code or status["unmatched"] or status["review"]:
        decide(f"**Codes for new facilities.** Every master-list facility needs a unique code "
               f"({without_code:,} registry rows have none). Who issues codes for facilities "
               f"marked ADD: the registry, or the portal?")
    if abs(len(reg.facilities) - EXPECTED_TOTAL) > EXPECTED_TOTAL * 0.05:
        decide(f"**Facility count.** The registry has {len(reg.facilities):,} Kaduna "
               f"facilities against the ≈{EXPECTED_TOTAL:,} expected. Which number is right, "
               f"and is the difference out-of-scope facility types?")

    paths = "\n".join(f"| {fld} | {', '.join(f'`{p}` ({c:,})' for p, c in cnt.most_common(4))} |"
                      for fld, cnt in ctx.sub_paths.items()) or "| — | — |"

    return f"""# Facility reconciliation

Generated {ctx.generated:%Y-%m-%d %H:%M} by `discovery/facility_reconcile.py`.
Workbook for client sign-off: `discovery/out/facility_reconciliation.xlsx`.

**Client decisions are needed by {ctx.deadline:%A %d %B %Y}** ({DECISION_DAYS} working days).
The signed-off list seeds the `facilities` table (build step 4) and generates the ODK cascade
file the Phase 2 form depends on (`ARCHITECTURE.md` §5.1), so it must be settled before
Phase 2 (Ingestion) starts.

## Inputs

| | |
|---|---|
| Registry file | `{ctx.registry_path.name}`: {reg.rows_read:,} rows read |
| Registry rows outside Kaduna State | {reg.other_state:,} {"" if "state" in reg.columns else "(no State column: the whole file was treated as Kaduna and filtered by LGA name)"} |
| Registry rows with an unrecognised LGA | {sum(reg.unknown_lga.values()):,} |
| Registry facilities kept (23 LGAs) | {len(reg.facilities):,} |
| Registry rows repeating a code | {reg.duplicate_codes:,} |
| Registry columns used | {", ".join(f"{k} = `{v}`" for k, v in reg.columns.items())} |
| Submissions source | `{ctx.submissions_path.name}`: {ctx.sub_records:,} records |
| Submissions without LGA and name/code | {ctx.sub_skipped:,} (skipped) |
| Distinct submitted facilities | {len(ctx.subs):,} |

Submission fields were read from these question paths (records per path):

| Field | Paths |
|---|---|
{paths}

## Method

1. **Normalise names:** lower case; strip punctuation and apostrophes; join spelled-out
   initials (`P.H.C.` → `phc`); expand or unify abbreviations (`Comp.` → comprehensive,
   `Center` → centre, `Clin.` → clinic, `PHCC` → `phc`); collapse facility-type phrases to one
   token: Primary Health (Care) Centre → `phc`, Comprehensive Health Centre → `chc`,
   Maternal (and) Child Health → `mch`, Health Centre/Center → `hc`, Health Post → `hp`.
   A value with no letters (e.g. `2019.00`) is not treated as a name.
2. **Group submissions** into distinct facilities by LGA plus registry code, or by LGA plus
   normalised name when the code is missing or not in the registry. A junk code shared by many
   facilities therefore never merges them.
3. **Match within the same LGA only**, in order: exact code (score 100); exact normalised
   name (100); fuzzy `token_set_ratio` on normalised names, with the score recorded. LGAs are
   resolved by exact spelling (plus known variants), never fuzzily.
4. **Sort by score:** ≥ {AUTO_MATCH} → Matched; {REVIEW_MIN}–{AUTO_MATCH - 1} → Review (top 3
   candidates); < {REVIEW_MIN} → Unmatched. A fuzzy score of {AUTO_MATCH}+ is still sent to
   Review when the runner-up is within {AMBIGUITY_MARGIN} points, or when the names agree only
   on facility-type words (`token_set_ratio` gives 100 to "PHC" vs "PHC Kawo"; the
   facility-type-stripped names must reach {CORE_MIN} on `token_sort_ratio`).

## Results

| Outcome | Facilities |
|---|---|
| Matched: by code | {methods['code']:,} |
| Matched: by exact name | {methods['exact name']:,} |
| Matched: fuzzy ≥ {AUTO_MATCH} | {methods['fuzzy']:,} |
| Review | {status['review']:,} |
| Unmatched | {status['unmatched']:,} |
| **Distinct submitted facilities** | **{len(ctx.subs):,}** |
| Registry facilities never assessed | {len(ctx.never_assessed):,} |

Why rows went to Review: {counts(review_reasons)}.
Why rows are Unmatched: {counts(unmatched_reasons)}.

### Per LGA

{md_table(LGA_HEADERS, lga_rows + [totals])}

### Public vs private (registry classification)

{md_table(["Ownership", "Registry", "Assessed (matched)", "Never assessed"], ownership_counts(ctx))}

### Totals

{md_table(["", "Count"], total_rows(ctx))}

Distinct submitted facilities can overstate the true number: spelling variants of one facility
that did not share a code count separately until the Review decisions merge them.

## Client decisions needed (by {ctx.deadline:%d %B %Y})

{chr(10).join(decisions)}
"""


# --------------------------------------------------------------------------- main


@dataclass
class Run:
    registry_path: Path
    submissions_path: Path
    registry: RegistryLoad
    subs: list[SubmittedFacility]
    sub_records: int
    sub_skipped: int
    sub_paths: dict[str, Counter]
    generated: datetime
    deadline: date
    assessed_registry: list[RegistryFacility] = field(default_factory=list)
    never_assessed: list[RegistryFacility] = field(default_factory=list)

    def __post_init__(self) -> None:
        hit = {id(s.match) for s in self.subs if s.status == "matched"}
        self.assessed_registry = [r for r in self.registry.facilities if id(r) in hit]
        self.never_assessed = [r for r in self.registry.facilities if id(r) not in hit]


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawTextHelpFormatter)
    p.add_argument("--registry", type=Path, required=True,
                   help="national health facility registry export (.csv or .xlsx)")
    p.add_argument("--registry-sheet", help="sheet name when the registry is .xlsx")
    p.add_argument("--map", action="append", default=[], metavar="FIELD=COLUMN",
                   help="registry column for FIELD (code, name, lga, ward, ownership, level, "
                        "state); repeatable")
    p.add_argument("--submissions", type=Path, default=HERE / "data" / "submissions",
                   help="OData JSON pages, or ODK CSV/XLSX exports: a file or a directory "
                        "(default: data/submissions/)")
    p.add_argument("--sub-map", action="append", default=[], metavar="FIELD=PATH",
                   help="submission question path for FIELD (lga, ward, code, name, ownership, "
                        "level), e.g. lga=grp_fac/lga; repeatable")
    p.add_argument("--lga-alias", action="append", default=[], metavar="VALUE=LGA",
                   help="treat VALUE (e.g. an ODK choice code) as one of the 23 LGAs; repeatable")
    p.add_argument("--deadline", type=date.fromisoformat,
                   help=f"decision deadline, YYYY-MM-DD (default: {DECISION_DAYS} working days)")
    return p.parse_args(argv)


def main(argv: list[str] | None = None) -> None:
    args = parse_args(argv)
    for path in (args.registry, args.submissions):
        if not path.exists():
            sys.exit(f"Not found: {path}")

    aliases: dict[str, str] = {}
    for pair in args.lga_alias:
        value, _, lga = pair.partition("=")
        canonical = resolve_lga(lga, {})
        if not value or canonical is None:
            sys.exit(f"--lga-alias {pair!r}: expected VALUE=LGA with LGA one of the 23")
        aliases[compact(value)] = canonical

    registry = load_registry(args.registry, args.registry_sheet,
                             parse_map(args.map, set(REGISTRY_FIELDS), "--map"), aliases)
    console.print("Registry columns: " + ", ".join(f"{k}={v!r}" for k, v in
                                                    registry.columns.items()))
    if "code" not in registry.columns:
        console.print("[yellow]No facility code column found: the exact-code step is skipped. "
                      "Pass --map code=... if the registry has one.[/]")
    if not registry.facilities:
        sys.exit("No registry facilities left in the 23 Kaduna LGAs; check --map and the file.")

    subs = load_submissions(args.submissions,
                            parse_map(args.sub_map, set(SUBMISSION_FIELDS), "--sub-map"),
                            aliases, {r.code for r in registry.facilities if r.code})
    match_all(subs.facilities, registry.facilities)

    generated = datetime.now()
    run = Run(
        registry_path=args.registry, submissions_path=args.submissions, registry=registry,
        subs=subs.facilities, sub_records=subs.records, sub_skipped=subs.skipped,
        sub_paths=subs.paths_used, generated=generated,
        deadline=args.deadline or add_business_days(generated.date(), DECISION_DAYS),
    )

    OUT_DIR.mkdir(exist_ok=True)
    build_workbook(run).save(XLSX_OUT)
    MD_OUT.write_text(build_markdown(run), encoding="utf-8")

    status = Counter(s.status for s in run.subs)
    table = Table(title="Facility reconciliation")
    table.add_column("Outcome")
    table.add_column("Count", justify="right")
    for label, value in (("Registry facilities (23 LGAs)", len(registry.facilities)),
                         ("Submitted facilities", len(run.subs)),
                         ("Matched", status["matched"]), ("Review", status["review"]),
                         ("Unmatched", status["unmatched"]),
                         ("Registry never assessed", len(run.never_assessed))):
        table.add_row(label, f"{value:,}")
    console.print(table)
    console.print(f"Wrote {XLSX_OUT.relative_to(HERE)} and {MD_OUT.relative_to(HERE)}")


if __name__ == "__main__":
    main()
